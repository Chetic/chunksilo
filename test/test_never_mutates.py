"""Regression guarantee: indexing NEVER writes inside an indexed directory.

The indexed directories are other people's documents; ChunkSilo must be a pure
reader of them. This test builds an index over a populated source tree and
proves the tree is byte-for-byte identical afterwards - no lock files, no
converted artifacts, no logs, no renames, nothing added or removed.
"""
import copy
import hashlib
import os
from pathlib import Path
from unittest.mock import patch

import pytest

from chunksilo import cfgload, index
from chunksilo.index import DirectoryConfig, IndexConfig, build_index


def _create_mock_embedding():
    from llama_index.core.embeddings import BaseEmbedding

    class MockEmbedding(BaseEmbedding):
        def _get_vector(self, text: str):
            vec = [0.0] * 384
            for word in text.lower().split():
                vec[sum(ord(c) for c in word) % 384] += 1.0
            return vec

        def _get_query_embedding(self, query: str):
            return self._get_vector(query)

        def _get_text_embedding(self, text: str):
            return self._get_vector(text)

        async def _aget_query_embedding(self, query: str):
            return self._get_vector(query)

    return MockEmbedding(model_name="mock-test")


def _populate(data_dir: Path) -> None:
    (data_dir / "notes.txt").write_text("plain text notes about the system")
    (data_dir / "guide.md").write_text("# Guide\n\nHow to configure the thing.\n")
    sub = data_dir / "sub dir with spaces"
    sub.mkdir()
    (sub / "deep.txt").write_text("content in a subdirectory")

    from docx import Document

    doc = Document()
    doc.add_paragraph("Heading text", style="Heading 1")
    doc.add_paragraph("Body of the docx document.")
    doc.save(str(data_dir / "report.docx"))

    # A legacy .doc as well: without LibreOffice it is skipped with a warning,
    # with LibreOffice it exercises the copy-first conversion path. Either way
    # the source tree must come through untouched.
    (data_dir / "legacy.doc").write_bytes(b"\xd0\xcf\x11\xe0 legacy doc bytes")


def _snapshot(root: Path) -> dict:
    """Every path under root, with content hash + size + mtime for files."""
    entries = {}
    for current, dirs, files in os.walk(root, followlinks=False):
        for name in dirs:
            path = Path(current) / name
            entries[str(path.relative_to(root))] = ("dir",)
        for name in files:
            path = Path(current) / name
            stat = path.lstat()
            entries[str(path.relative_to(root))] = (
                "file",
                stat.st_size,
                stat.st_mtime_ns,
                hashlib.sha256(path.read_bytes()).hexdigest(),
            )
    return entries


@pytest.fixture
def indexing_env(tmp_path, monkeypatch):
    data_dir = tmp_path / "source"
    data_dir.mkdir()
    storage_dir = tmp_path / "storage"

    # Run from a neutral CWD: anything that writes a relative path (a log, a
    # scratch file) must land here or in storage, never in the source tree.
    cwd = tmp_path / "cwd"
    cwd.mkdir()
    monkeypatch.chdir(cwd)

    test_index_config = IndexConfig(
        directories=[DirectoryConfig(path=data_dir)], chunk_size=512, chunk_overlap=100
    )

    # The active config governs cfgload.get() consumers (doc_temp location).
    active = copy.deepcopy(cfgload._DEFAULTS)
    active["storage"]["storage_dir"] = str(storage_dir)
    monkeypatch.setattr(cfgload, "_config_cache", active)
    monkeypatch.setattr(cfgload, "_active_path", tmp_path / "nonexistent.yaml")

    orig = (
        index.STORAGE_DIR,
        index.STATE_DB_PATH,
        index.BM25_INDEX_DIR,
        index.HEADING_STORE_PATH,
        index._config,
    )
    index.STORAGE_DIR = storage_dir
    index.STATE_DB_PATH = storage_dir / "ingestion_state.db"
    index.BM25_INDEX_DIR = storage_dir / "bm25_index"
    index.HEADING_STORE_PATH = storage_dir / "heading_store.json"
    index._heading_store = None

    mock_embed = _create_mock_embedding()
    with patch("chunksilo.index.load_index_config", return_value=test_index_config), \
         patch("chunksilo.index._create_fastembed_embedding", return_value=mock_embed), \
         patch("chunksilo.index.ensure_embedding_model_cached"), \
         patch("chunksilo.index.ensure_rerank_model_cached"):
        yield data_dir

    (
        index.STORAGE_DIR,
        index.STATE_DB_PATH,
        index.BM25_INDEX_DIR,
        index.HEADING_STORE_PATH,
        index._config,
    ) = orig
    index._heading_store = None


def test_indexing_never_mutates_the_source_tree(indexing_env):
    data_dir = indexing_env
    _populate(data_dir)

    before = _snapshot(data_dir)
    assert before, "populate created nothing?"

    build_index()          # initial build
    build_index()          # incremental rebuild over the same content

    after = _snapshot(data_dir)

    assert set(after) == set(before), (
        "indexing added or removed entries in the source tree: "
        f"added={set(after) - set(before)}, removed={set(before) - set(after)}"
    )
    for rel, entry in before.items():
        assert after[rel] == entry, f"indexing modified {rel}: {entry} -> {after[rel]}"

    # And nothing wrote a log into the tree under any name.
    assert not list(data_dir.rglob("*.log"))
    assert not list(data_dir.rglob(".~lock*"))
